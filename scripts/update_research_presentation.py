#!/usr/bin/env python3
"""Inspect and update the research presentation without replacing the original."""

import argparse
from collections import Counter
import hashlib
import json
from pathlib import Path
import subprocess
import xml.etree.ElementTree as ET
import zipfile


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs/Amharic_ACOS_Research_Presentation(v3.0).pptx"
OUTPUT = ROOT / "docs/Amharic_ACOS_Research_Presentation(v4.0).pptx"
DRAWING_NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
BLUE = "104563"
TEAL = "167D97"
ORANGE = "B85F24"
INK = "202B38"
MUTED = "536777"
LIGHT = "EAF2F6"
PALE = "FCF0E7"


def read_json(relative_path):
    return json.loads((ROOT / relative_path).read_text(encoding="utf-8-sig"))


def percent(value):
    return f"{value * 100:.2f}%"


def dataset_summary(directory):
    summaries = {}
    for split in ("train", "dev", "test"):
        with (ROOT / "data" / directory / f"{split}.jsonl").open(encoding="utf-8") as handle:
            records = [json.loads(line) for line in handle]
        quads = [quad for record in records for quad in record["quads"]]
        summaries[split] = {
            "sentences": len(records),
            "quads": len(quads),
            "implicit": sum(quad["a_start"] < 0 or quad["o_start"] < 0 for quad in quads),
            "categories": Counter(quad["category"] for quad in quads),
        }
    return summaries


class Deck:
    def __init__(self, source):
        from pptx import Presentation

        self.presentation = Presentation(source)
        self.entries = []
        self.commit = subprocess.check_output(["git", "rev-parse", "HEAD"], cwd=ROOT, text=True).strip()
        self.tracked = set(subprocess.check_output(["git", "ls-files", "-z"], cwd=ROOT, text=True).split("\0"))
        self.references = self.presentation.slides[-1]
        for slide_id in list(self.presentation.slides._sldIdLst)[1:-1]:
            self.presentation.part.drop_rel(slide_id.rId)
            self.presentation.slides._sldIdLst.remove(slide_id)
        from pptx.opc.packuri import PackURI

        self.references.part.partname = PackURI("/ppt/slides/slide2.xml")
        cover = self.presentation.slides[0]
        for shape in cover.shapes:
            if not shape.has_text_frame:
                continue
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.text = run.text.replace(
                        "Amharic E-Commerce & Hospitality Reviews", "Amharic Civic & Governance Commentary"
                    ).replace(
                        "ByT5-Base Generative (Final Test Evaluation)", "ByT5 Run 1 Results + Implementation & Data Audit"
                    ).replace("September 2026", "October 2026")
        self.text(cover, "v4.0 | Repository evidence update: 2 October 2026", 0.7, 6.45, 11.5, 0.3, 13, "A6CADB")
        self.record(cover, "Research scope and current status", ["data/README.md", "results/stage_byt5_acos/byt5_base_run1/test_metrics.json"])

    def text(self, slide, content, left, top, width, height, size=18, color=INK, bold=False, font="Calibri"):
        from pptx.dml.color import RGBColor
        from pptx.util import Inches, Pt

        shape = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
        frame = shape.text_frame
        frame.word_wrap = True
        frame.margin_left = frame.margin_right = 0
        frame.margin_top = frame.margin_bottom = 0
        for index, line in enumerate(content.split("\n")):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.text = line
            paragraph.font.name = font
            paragraph.font.size = Pt(size)
            paragraph.font.bold = bold
            paragraph.font.color.rgb = RGBColor.from_string(color)
            paragraph.space_after = Pt(8)
        return shape

    def block(self, slide, left, top, width, height, color=LIGHT):
        from pptx.dml.color import RGBColor
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.util import Inches

        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(color)
        shape.line.fill.background()
        return shape

    def record(self, slide, title, sources, note=""):
        source_lines = []
        for source in sources:
            path = source.split("#", 1)[0]
            if source.startswith("https://"):
                source_lines.append(source)
            elif path in self.tracked or (ROOT / path).is_dir():
                kind = "tree" if (ROOT / path).is_dir() else "blob"
                source_lines.append(f"https://github.com/EthAIResearch/acos-amharic/{kind}/{self.commit}/{source}")
            else:
                source_lines.append(f"Local working-tree file (not committed at audit): {source}")
        slide.notes_slide.notes_text_frame.text = (
            f"Evidence audit: 2026-10-02\n{title}\n\n{note}\n\nSources:\n" + "\n".join(source_lines)
        )
        self.entries.append({"title": title, "sources": sources, "note": note})

    def slide(self, section, title, subtitle, sources, note=""):
        slide = self.presentation.slides.add_slide(self.references.slide_layout)
        for placeholder in list(slide.placeholders):
            placeholder._element.getparent().remove(placeholder._element)
        self.text(slide, section.upper(), 0.55, 0.30, 12.2, 0.25, 11, TEAL, True)
        self.text(slide, title, 0.55, 0.68, 12.2, 0.65, 28, INK, True, "Cambria")
        self.text(slide, subtitle, 0.55, 1.36, 12.2, 0.50, 15, MUTED)
        self.block(slide, 0.55, 6.98, 12.2, 0.015, "CCDCE4")
        short_sources = " | ".join(
            "GitHub #" + source.rsplit("/", 1)[-1] if source.startswith("https://") else source
            for source in sources
        )
        if len(short_sources) > 178:
            short_sources = short_sources[:155] + " ... (full sources in notes)"
        self.text(slide, short_sources, 0.55, 7.10, 11.9, 0.22, 8, MUTED)
        self.record(slide, title, sources, note)
        return slide

    def table(self, slide, headers, rows, top=2.0, widths=None, left=0.55, row_height=0.55, size=17):
        from pptx.dml.color import RGBColor
        from pptx.enum.text import MSO_ANCHOR
        from pptx.util import Inches, Pt

        widths = widths or [12.2 / len(headers)] * len(headers)
        table = slide.shapes.add_table(
            len(rows) + 1, len(headers), Inches(left), Inches(top), Inches(sum(widths)),
            Inches(row_height * (len(rows) + 1)),
        ).table
        for column, width in zip(table.columns, widths):
            column.width = Inches(width)
        for row_index, values in enumerate([headers, *rows]):
            table.rows[row_index].height = Inches(row_height)
            for column_index, value in enumerate(values):
                cell = table.cell(row_index, column_index)
                cell.text = str(value)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                cell.margin_left = Inches(0.12)
                cell.margin_right = Inches(0.10)
                cell.margin_top = cell.margin_bottom = Inches(0.04)
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor.from_string(BLUE if row_index == 0 else (LIGHT if row_index % 2 else "FFFFFF"))
                for paragraph in cell.text_frame.paragraphs:
                    paragraph.font.name = "Calibri"
                    paragraph.font.size = Pt(size)
                    paragraph.font.bold = row_index == 0
                    paragraph.font.color.rgb = RGBColor.from_string("FFFFFF" if row_index == 0 else INK)
        return table

    def columns(self, slide, left_title, left_lines, right_title, right_lines, size=19):
        for left, title, lines, accent in (
            (0.55, left_title, left_lines, BLUE), (6.85, right_title, right_lines, ORANGE)
        ):
            self.block(slide, left, 2.02, 5.9, 0.04, accent)
            self.text(slide, title, left, 2.20, 5.9, 0.56, 21, accent, True, "Cambria")
            self.text(slide, "\n".join(lines), left, 2.91, 5.9, 3.75, size)

    def takeaway(self, slide, content, top=5.85):
        self.block(slide, 0.55, top, 12.2, 0.82, LIGHT)
        self.text(slide, content, 0.75, top + 0.14, 11.8, 0.59, 18, BLUE)

    def flow(self, slide, labels, descriptions, top=2.35):
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.dml.color import RGBColor
        from pptx.util import Inches

        count = len(labels)
        gap = 0.35
        width = (12.2 - gap * (count - 1)) / count
        for index, (label, description) in enumerate(zip(labels, descriptions)):
            left = 0.55 + index * (width + gap)
            self.block(slide, left, top, width, 1.0, BLUE if index % 2 == 0 else TEAL)
            self.text(slide, label, left + 0.12, top + 0.15, width - 0.24, 0.76, 18, "FFFFFF", True)
            self.text(slide, description, left, top + 1.22, width, 1.40, 16, MUTED)
            if index < count - 1:
                arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(left + width + 0.045), Inches(top + 0.4), Inches(0.26), Inches(0.18))
                arrow.fill.solid()
                arrow.fill.fore_color.rgb = RGBColor.from_string(TEAL)
                arrow.line.fill.background()

    def save(self, output):
        identifiers = self.presentation.slides._sldIdLst
        reference_id = identifiers[1]
        identifiers.remove(reference_id)
        identifiers.append(reference_id)
        self.record(self.references, "Bibliography (retained from v3.0)", ["docs/references.md"])
        for number, slide in enumerate(self.presentation.slides, 1):
            for shape in slide.shapes:
                if shape.has_text_frame and shape.text.strip().isdigit() and shape.top > 6.9 * 914400:
                    for paragraph in shape.text_frame.paragraphs:
                        for run in paragraph.runs:
                            run.text = str(number)
            if number != len(self.presentation.slides):
                self.text(slide, str(number), 12.85, 7.1, 0.38, 0.22, 9, MUTED)
        output.parent.mkdir(parents=True, exist_ok=True)
        self.presentation.save(output)
        with zipfile.ZipFile(output) as archive:
            assert len(archive.namelist()) == len(set(archive.namelist())), "Duplicate package part names"


def github(number, kind="pull"):
    return f"https://github.com/EthAIResearch/acos-amharic/{kind}/{number}"


def add_problem_slide(deck):
    slide = deck.slide(
        "Problem formulation", "Amharic ACOS: Four Fields, One Exact Match",
        "Output a set of aspect-category-sentiment-opinion tuples for each civic-domain sentence.",
        ["src/common/data_prep.py", "src/stage_byt5_acos/linearization.py", "src/pipeline/evaluate.py"]
    )
    deck.block(slide, 0.55, 2.0, 12.2, 0.95, LIGHT)
    deck.text(slide, "[ aspect | category | sentiment | opinion ] ; ...", 0.8, 2.23, 11.7, 0.48, 25, BLUE, True)
    deck.text(slide, "NULL = an implicit aspect or opinion. None = a sentence with no quadruples.", 0.55, 3.15, 12.2, 0.5, 21)
    deck.table(slide, ["Challenge", "Evidence", "Modeling consequence"], [
        ["Implicit components", "42.71% of training quads", "Need NULL prediction or implicit classifiers"],
        ["Unseen opinion phrases", "1,969 / 2,983 test opinions", "Phrase generalization, not just vocabulary coverage"],
        ["Rich morphology", "Clitics and multi-word opinions", "Strict boundaries remain difficult"]
    ], top=3.75, widths=[3.0, 3.6, 5.6], row_height=0.55, size=16)
    deck.takeaway(slide, "Unseen surface phrases are not the same as tokenizer <unk> tokens. Byte coverage does not guarantee semantic accuracy.", top=6.10)


def add_pipeline_slides(deck):
    slide = deck.slide(
        "Phase 1 | implemented pipeline", "Six Stages, Including Learned Implicit Handling",
        "Stage 6 consumes predictions from earlier stages; gold-span component scores are diagnostics only.",
        ["src/pipeline/run_inference.py", "src/pipeline/assemble_quads.py", github(47)]
    )
    deck.flow(slide, ["1. Spans", "2. Pairs", "3. Category", "4. Sentiment", "5. Implicit", "6. Assemble"], [
        "Shared encoder; dual BIO heads", "Union of nearest-distance pairs", "Span + sentence pooling; category head",
        "Shared PairClassifier; 3 polarities", "Anchored / sentence-level classifiers", "Full, explicit and implicit exact-match F1"
    ])
    deck.takeaway(slide, "Missing Stage 5 category/sentiment models leave a recall gap. This architecture is not inherently unable to emit NULL.")

    slide = deck.slide(
        "Baseline comparison | Stages 1, 3 and 4", "Backbone Benchmarks Across Three Subtasks",
        "Recorded component results: Stage 3/4 use gold explicit pairs, not noisy end-to-end predictions.",
        ["results/stage1/", "results/stage3/", "results/stage4/", github(36), github(33)]
    )
    rows = []
    for name, run in (("BERT-small", "bertsmall_base_run1"), ("AfroXLMR-base", "afroxlmr_base_run1"), ("RoBERTa-Amharic", "roberta_base_run1")):
        category = read_json(f"results/stage3/{run}/best_metrics.json")
        sentiment = read_json(f"results/stage4/{run}/best_metrics.json")
        stage1_path = ROOT / f"results/stage1/{run}/best_metrics.json"
        extraction = read_json(str(stage1_path.relative_to(ROOT))) if stage1_path.exists() else None
        rows.append([name, percent(extraction["aspect"]["f1"]) if extraction else "Not recorded",
                     percent(extraction["opinion"]["f1"]) if extraction else "Not recorded",
                     percent(category["macro_f1_present_categories"]), percent(sentiment["macro_f1"]), percent(sentiment["accuracy"])])
    deck.table(slide, ["Backbone", "Aspect F1", "Opinion F1", "Category\nmacro-F1*", "Sentiment\nmacro-F1", "Sentiment\naccuracy"], rows,
               widths=[2.6, 1.9, 1.9, 2.0, 1.9, 1.9], row_height=0.8, size=16)
    deck.text(slide, "*Present-label category macro-F1. RoBERTa: 62.25% present-label vs 36.79% all-label macro-F1.\nNine declared categories have no evaluation support. RoBERTa neutral-class F1 is only 33.78%.", 0.55, 5.45, 12.2, 1.0, 18)

    slide = deck.slide(
        "Implicit modeling | Issues #30, #38-44", "Detection Is Not Complete Quadruple Recovery",
        "Stage 5 reuses PairClassifier; absent-span masks become zero vectors while sentence context remains available.",
        ["docs/stage5_implicit_analysis.md", "results/stage5/", "configs/stage6_pipeline.yaml", github(34)]
    )
    rows = []
    for label, run, positive_f1 in (
        ("Implicit aspect / opinion anchor", "detect_implicit_aspect_run1", "63.81%"),
        ("Implicit opinion / aspect anchor", "detect_implicit_opinion_run1", "47.05%"),
        ("Both implicit / sentence anchor", "detect_fully_implicit_run1", "35.17%"),
        ("Category / implicit aspect", "category_opinion_anchor_run1", "Not a detector"),
    ):
        metrics = read_json(f"results/stage5/{run}/best_metrics.json")
        rows.append([label, percent(metrics["accuracy"]), percent(metrics["macro_f1_present_categories"]), positive_f1])
    deck.table(slide, ["Subtask", "Accuracy", "Macro-F1*", "Positive-class F1"], rows,
               widths=[5.3, 2.0, 2.0, 2.9], row_height=0.64, size=17)
    deck.text(slide, "Nine trainable sub-models; four recorded result sets. *Category row: present-label macro-F1.\nFive completion-model result sets are absent; fully implicit training keeps only the first such quad per sentence.", 0.55, 5.5, 12.2, 1.0, 18)

    slide = deck.slide(
        "Phase 1 diagnosis", "Error Propagation: Measured, Not an F1 Product",
        "Issue #51 reports approximately 14% full-pipeline F1; no standalone Stage 6 metric JSON is present.",
        [github(51, "issues"), "docs/stage2_pairing_analysis.md", "src/pipeline/run_inference.py"],
        "The original deck's 14.2% is a historical narrative figure, not independently reproduced here. Multiplying component F1 scores is not an end-to-end recall derivation."
    )
    deck.columns(slide, "What the evidence says", [
        "Nearest-distance pairing reaches about 99.3% F1 when both input spans are gold.",
        "Stage 1 misses and boundary shifts change the candidate pairs available downstream.",
        "A later classifier cannot repair a missing explicit candidate in that branch."
    ], "What to report", [
        "Full-system P/R/F1 from predictions passed through every stage.",
        "Separate component diagnostics, plus explicit and implicit quad subsets.",
        "Do not multiply independent F1 scores or compare gold-pair accuracy with full-quad F1."
    ])


def add_sdrn_slides(deck):
    slide = deck.slide(
        "Phase 2 | joint extraction", "SDRN: Entity and Relation Feedback",
        "The implementation evolved from dual-head approximation to unified 5-way CRF tagging and word-level pooling.",
        ["src/stage_aope_sdrn/model.py", "src/stage_aope_sdrn/crf.py", github(64), github(75)]
    )
    deck.flow(slide, ["Context encoder", "Entity channel", "Relation channel", "Joint pairs"], [
        "AfroXLMR; optional subword-to-word mean pooling", "O / B-ASP / I-ASP / B-OPN / I-OPN; constrained CRF",
        "Row-normalized attention; ESM and RSM recurrent feedback", "Bidirectional span correlation; acceptance threshold"
    ])
    deck.takeaway(slide, "The final word-level model tags whole words. Its remaining errors cannot be explained solely by subword fragmentation.")

    slide = deck.slide(
        "Ablation history | Issues #58, #69, #76-79", "What Helped, and What Did Not",
        "Recorded test metric artifacts; runs change more than one factor and are not controlled single-factor ablations.",
        ["results/stage_aope_sdrn/", github(70), github(78), github(80), github(76, "issues")],
        "CRF Run 1 best_metrics uses threshold 0.20; other rows are recorded defaults around 0.50. Do not interpret these as identical operating points."
    )
    rows = []
    for label, run in (
        ("Explicit data + weights", "afroxlmr_base_run1"), ("Exact SDRN + CRF", "afroxlmr_crf_run1"),
        ("Asymmetric weights", "afroxlmr_crf_run2"), ("+ active Dice loss", "afroxlmr_crf_run3"),
        ("Word-level SDRN", "afroxlmr_word_sdrn_run1"),
    ):
        metrics = read_json(f"results/stage_aope_sdrn/{run}/best_metrics.json")
        rows.append([label, percent(metrics["aspect"]["f1"]), percent(metrics["opinion"]["recall"]), percent(metrics["candidate_ceiling_recall"]), percent(metrics["f1"])])
    deck.table(slide, ["Run / change", "Aspect F1", "Opinion recall", "Pair ceiling", "Pair F1"], rows,
               widths=[3.7, 2.0, 2.3, 2.1, 2.1], row_height=0.59, size=17)
    deck.takeaway(slide, "Dice loss did not improve pair F1 in the recorded run. Thresholds differ: CRF Run 1 uses 0.20; most others use 0.50.", top=6.05)

    slide = deck.slide(
        "Word-level SDRN | diagnostic evidence", "Seen vs. Unseen: Token Overlap & Span Boundary Breakdown",
        "Recall across 3 granularities: full phrase, >=1 token overlap, and character trigrams for both aspect and opinion.",
        ["results/stage_aope_sdrn/afroxlmr_word_sdrn_run1/seen_unseen_diagnostic.json", github(76, "issues")]
    )
    deck.table(slide, ["Span Type & Granularity", "Seen in Train", "Seen Recall", "Unseen in Train", "Unseen Recall"], [
        ["Opinion: Exact full phrase", "1,014 (34.0%)", "59.17% (600)", "1,969 (66.0%)", "37.63% (741)"],
        ["Opinion: >=1 token overlap", "2,169 (72.7%)", "45.78% (993)", "814 (27.3%)", "42.75% (348)"],
        ["Opinion: Character trigram >=50%", "1,450 (48.6%)", "54.41% (789)", "1,533 (51.4%)", "36.01% (552)"],
        ["Opinion: Multi-word phrase", "213 (15.2%)", "62.44% (133)", "1,184 (84.8%)", "31.93% (378)"],
        ["Aspect: Exact full phrase", "1,910 (64.0%)", "66.65% (1,273)", "1,073 (36.0%)", "43.80% (470)"],
        ["Aspect: >=1 token overlap", "2,550 (85.5%)", "59.73% (1,523)", "433 (14.5%)", "50.81% (220)"],
        ["Aspect: Character trigram >=50%", "2,326 (78.0%)", "63.67% (1,481)", "657 (22.0%)", "39.88% (262)"],
    ], top=1.92, widths=[3.6, 2.1, 2.1, 2.2, 2.2], row_height=0.45, size=14)
    deck.takeaway(slide, "72.7% of opinions share >=1 token with train, but exact boundary recall collapses on multi-word unseen opinions (31.93%). SDRN errors stem from token-level boundary fragmentation, not pure lexical absence.", top=5.75)

    slide = deck.slide(
        "Transition to span modeling", "Candidate Coverage Is Not Final Pair Recall",
        "Denominators matter: reachable gold pairs / all gold pairs versus accepted gold pairs / reachable gold pairs.",
        ["results/stage_aope_sdrn/afroxlmr_crf_run1/best_metrics.json", "results/stage_span_aste/afroxlmr_pooled_focal_run3/best_metrics.json"]
    )
    deck.table(slide, ["Model", "Reachable / gold", "Candidate recall", "Accepted / reachable"], [
        ["SDRN CRF Run 1", "979 / 2,963", "33.04%", "966 / 979 = 98.67%"],
        ["Span-ASTE Run 1", "2,403 / 2,963", "81.10%", "706 / 2,403 = 29.38%"],
        ["Span-ASTE Run 3", "2,427 / 2,963", "81.91%", "1,093 / 2,427 = 45.04%"],
    ], widths=[3.4, 2.8, 2.5, 3.5], row_height=0.80, size=17)
    deck.takeaway(slide, "Span enumeration raises reachability; relation classification then becomes the larger bottleneck. These are run-specific observations.")


def add_dataset_slides(deck):
    source = dataset_summary("prepared")
    clean = dataset_summary("prepared_clean")
    slide = deck.slide(
        "Dataset and annotation", "A Civic-Domain Dataset, Not a Review Benchmark",
        "Source text: mT5/mC4/OSCAR; Gemini annotations as described in the repository.",
        ["data/README.md", "data/prepared/train.jsonl", "data/prepared/test.jsonl"],
        "Counts recomputed from the JSONL files. No independent human annotation-quality study is recorded."
    )
    deck.table(slide, ["Original split", "Sentences", "Quadruples", "Implicit quads"], [
        [split.title(), f"{values['sentences']:,}", f"{values['quads']:,}", percent(values["implicit"] / values["quads"])]
        for split, values in source.items()
    ], widths=[3.5, 2.7, 2.7, 3.3], row_height=0.7)
    deck.text(slide, "Explicit-only training: 22,765 sentences / 23,617 quads.\nImplicitness is split-specific: 42.71% in train, 41.94% in the original test set.", 0.55, 5.10, 12.2, 0.72, 19)
    deck.takeaway(slide, "LLM annotation is a provenance fact, not a quality guarantee. Human adjudication remains a research requirement.")

    slide = deck.slide(
        "Data quality | Issue #103, PR #104", "A Cleaned Dataset Is Now Available",
        "The cleaning pass changes both training data and evaluation annotations; old scores are historical baselines.",
        ["data/clean_dataset.py", "data/prepared_clean/cleaning_report.json", "https://github.com/EthAIResearch/acos-amharic/pull/104"],
        "Independent checks found zero normalized-text duplicates, split overlap, duplicate quads, and same-span sentiment conflicts in prepared_clean. Sentiment normalization is an automatic heuristic, not human adjudication."
    )
    deck.table(slide, ["Clean split", "Sentences", "Quadruples", "Implicit quads"], [
        [split.title(), f"{values['sentences']:,}", f"{values['quads']:,}", percent(values["implicit"] / values["quads"])]
        for split, values in clean.items()
    ], widths=[3.5, 2.7, 2.7, 3.3], row_height=0.61)
    deck.text(slide, "Merged duplicate rows: 233 train / 5 dev / 1 test. Removed leaked rows: 116 train / 10 dev.\nDropped 527 training rows over 100 tokens; resolved sentiment conflicts using train-majority labels.\nOriginal overlap: 68 train-dev / 48 train-test / 10 dev-test texts. Clean overlap: zero.", 0.55, 4.80, 12.2, 1.2, 18)
    deck.takeaway(slide, "No clean-data model metrics are recorded. Retrain and evaluate all compared models on the same data version.", top=6.05)

    slide = deck.slide(
        "Schema and long-tail coverage", "22 Declared Categories; 23 Observed Labels",
        "The taxonomy is externally fixed. Cleaning preserves labels; it does not reconcile the schema.",
        ["data/prepared/label_space.json", "src/stage_byt5_acos/linearization.py", "data/prepared_clean/test.jsonl"],
        "ECONOMY#UTILITIES occurs once in original and clean test data, never in train/dev, and is absent from CATEGORIES and label_space.json. Older README says 13; v3.0 says 29. Neither matches current declared metadata."
    )
    deck.columns(slide, "Observed contract", [
        "Seven civic domains; 22 labels in label_space.json and ByT5 CATEGORIES.",
        "A 23rd label, ECONOMY#UTILITIES, occurs once in test and has no training support.",
        "The strict ByT5 parser rejects that undeclared category."
    ], "Reporting implications", [
        "Keep the category names unchanged; reconcile the allowed list with the data owner.",
        "Nine categories are absent from the recorded Stage 3 evaluation.",
        "Report macro-F1 over present labels AND over all declared labels."
    ])


def add_example_slide(deck):
    with (ROOT / "data/prepared/test.jsonl").open(encoding="utf-8") as handle:
        records = [json.loads(line) for line in handle]
    slide = deck.slide(
        "Worked Amharic examples", "Explicit and NULL Fields in the Actual Dataset",
        "Original test annotations, not model predictions or independently adjudicated reference labels.",
        ["data/prepared/test.jsonl#L3367", "data/prepared/test.jsonl#L151", "src/stage_byt5_acos/linearization.py"],
        "Examples are read directly from 1-based rows 3367 and 151. The implicit example has a category annotation without an explicit educational target, illustrating why human adjudication is important."
    )
    for left, number, title in ((0.55, 3367, "Explicit aspect and opinion"), (6.85, 151, "Implicit aspect; explicit opinion")):
        record = records[number - 1]
        quad = record["quads"][0]
        tokens = record["tokens"]
        assert quad["category"] == "PUBLIC_SERVICES#EDUCATION" and quad["sentiment"] == "POSITIVE"
        aspect = " ".join(tokens[quad["a_start"]:quad["a_end"]]) if quad["a_start"] >= 0 else "NULL"
        opinion = " ".join(tokens[quad["o_start"]:quad["o_end"]]) if quad["o_start"] >= 0 else "NULL"
        deck.block(slide, left, 2.02, 5.9, 0.04, TEAL)
        deck.text(slide, title, left, 2.20, 5.9, 0.45, 21, BLUE, True, "Cambria")
        deck.text(slide, record["text"], left, 2.93, 5.9, 0.85, 25, INK, font="Droid Sans Ethiopic")
        deck.text(slide, f"Aspect: {aspect}\nOpinion: {opinion}", left, 3.96, 5.9, 1.15, 20, TEAL, font="Droid Sans Ethiopic")
        deck.text(slide, f"Category: {quad['category']}\nSentiment: {quad['sentiment']}  |  Test row: {number}", left, 5.12, 5.9, 0.75, 16)
    deck.takeaway(slide, "An implicit category is an annotation decision, not a recoverable span. These examples motivate human review of contextual labels.", top=6.10)


def add_span_slides(deck):
    run_names = ["afroxlmr_run1", "afroxlmr_biaffine_run2", "afroxlmr_pooled_focal_run3"]
    metrics = [read_json(f"results/stage_span_aste/{run}/best_metrics.json") for run in run_names]
    slide = deck.slide(
        "Phases 3-5 | span paradigm", "Span-ASTE: Enumerate, Prune, Then Classify",
        "Whole-span candidates replace BIO decisions; this implementation extracts explicit pairs and sentiment triplets.",
        ["src/stage_span_aste/model.py", "src/stage_span_aste/span_utils.py", github(82), github(88)]
    )
    deck.flow(slide, ["Word encoding", "Span features", "Top-k pruning", "Pair scoring", "Output"], [
        "AfroXLMR subwords pooled to words", "All contiguous spans up to 8 words; learned width features",
        "Separate aspect and opinion mention scores", "MLP, then biaffine scoring + cross-span context",
        "INVALID / POS / NEG / NEU; no category head or NULL candidate"
    ])
    deck.takeaway(slide, "Candidate spans are bounded by encoded words to avoid truncation-related indexing failures (PR #83).")

    slide = deck.slide(
        "Span-ASTE Run 1 | recorded test results", "High Candidate Coverage, Low Conversion",
        "Run 1 reaches 2,403 of 2,963 gold pairs, but accepts only 706 correct pairs at its recorded operating point.",
        ["results/stage_span_aste/afroxlmr_run1/best_metrics.json", github(84, "issues"), github(86)]
    )
    deck.table(slide, ["Task", "Precision", "Recall", "F1", "True positives"], [
        [label, percent(metrics[0][key]["precision"]), percent(metrics[0][key]["recall"]), percent(metrics[0][key]["f1"]), metrics[0][key]["tp"]]
        for label, key in (("Aspect extraction", "ate"), ("Opinion extraction", "ote"), ("AOPE pairs", "aope"), ("ASTE triplets", "aste"))
    ], widths=[4.0, 2.0, 2.0, 2.0, 2.2], row_height=0.64)
    deck.takeaway(slide, "Only 29.38% of reachable gold pairs become true positives. Threshold sweeps diagnose conservatism; tune on dev only.")

    slide = deck.slide(
        "Span-ASTE Run 2 | biaffine interaction", "A Better Relation Model, With Modest Gains",
        "Deep projections, bilinear compatibility and cross-span attention replace the concatenation-only MLP.",
        ["results/stage_span_aste/afroxlmr_biaffine_run2/best_metrics.json", "src/stage_span_aste/model.py", github(91)]
    )
    rows = []
    for label, key in (("Aspect F1", "ate"), ("Opinion F1", "ote"), ("AOPE pair F1", "aope"), ("ASTE triplet F1", "aste")):
        before, after = metrics[0][key]["f1"], metrics[1][key]["f1"]
        rows.append([label, percent(before), percent(after), f"{100 * (after-before):+.2f} pts"])
    deck.table(slide, ["Recorded metric", "Run 1", "Run 2", "Change"], rows,
               widths=[5.0, 2.4, 2.4, 2.4], row_height=0.64)
    deck.takeaway(slide, "34.02% is Run 2's candidate conversion, not its AOPE F1. Contextual endpoints are not a direct pool of internal words.")

    slide = deck.slide(
        "Span-ASTE Run 3 | implementation details", "Full-Span Pooling and Cost-Sensitive Loss",
        "The result artifact records weighted cross-entropy. Focal loss exists in code but was OFF for this run.",
        ["src/stage_span_aste/model.py", "results/stage_span_aste/afroxlmr_pooled_focal_run3/resolved_args.json", github(93)]
    )
    deck.columns(slide, "Constant-time span queries", [
        "Use half-open spans [s, e), matching the annotation convention.",
        "P[0] = 0; P[t] = sum of x[0:t].\nmean(s,e) = (P[e] - P[s]) / (e - s).",
        "h_span = [x_s; x_(e-1); mean; width].\nDimension: 3 x 768 + 20 = 2,324.",
        "O(n) prefix preprocessing; O(1) per span for fixed hidden width."
    ], "Recorded loss configuration", [
        "Weights [INVALID, POS, NEG, NEU] = [0.1, 1.0, 1.0, 2.0].",
        "use_span_mean_pooling = true\nuse_focal_loss = false",
        "Optional focal form in code:\n-alpha * (1 - p)^gamma * log(p).",
        "Pooling and loss weights changed together; no isolated causal ablation."
    ], size=18)

    slide = deck.slide(
        "Span-ASTE Run 3 | recall trade-off", "More Opinions Found, Not a Universal F1 Gain",
        "Default metric artifacts are shown together; separate sweep outputs are reported on the next slide.",
        ["results/stage_span_aste/afroxlmr_run1/best_metrics.json", "results/stage_span_aste/afroxlmr_pooled_focal_run3/best_metrics.json"]
    )
    rows = []
    for label, key, field in (("Aspect F1", "ate", "f1"), ("Opinion F1", "ote", "f1"), ("Opinion recall", "ote", "recall"), ("AOPE pair F1", "aope", "f1"), ("AOPE pair recall", "aope", "recall")):
        rows.append([label, *[percent(result[key][field]) for result in metrics]])
    deck.table(slide, ["Metric", "Run 1", "Run 2", "Run 3"], rows,
               widths=[5.0, 2.4, 2.4, 2.4], row_height=0.59)
    deck.takeaway(slide, "55.20% aspect F1 is best among these Span-ASTE runs, not across the repository: SDRN CRF Run 1 records 60.41%.", top=6.05)

    slide = deck.slide(
        "Calibration and experiment integrity", "Dev-Tuned Results Are Not Test-Selected Peaks",
        "Record the metric file, data version, threshold-selection split and decoding rule for every reported score.",
        ["results/stage_span_aste/afroxlmr_biaffine_run2/dev_threshold_sweep.json", "results/stage_span_aste/afroxlmr_pooled_focal_run3/threshold_sweep.json"],
        "best_metrics and threshold_sweep do not exactly reproduce each other at 0.50. Keep their provenance separate; do not mix counts or derive gains across files without rerunning a frozen evaluator."
    )
    rows = []
    for label, run in (("Biaffine Run 2", run_names[1]), ("Pooled Run 3", run_names[2])):
        dev = read_json(f"results/stage_span_aste/{run}/dev_threshold_sweep.json")
        test = read_json(f"results/stage_span_aste/{run}/threshold_sweep.json")
        threshold = dev["best_aope_threshold"]
        dev_tuned = next(row for row in test["threshold_sweep"] if row["threshold"] == threshold)
        rows.append([label, f"{threshold:.2f}", percent(dev_tuned["aope_f1"]), f"{test['best_aope_threshold']:.2f}", percent(test["best_aope_f1"])])
    deck.table(slide, ["Run", "Dev-selected\nthreshold", "Test F1 at\ndev threshold", "Test-selected\nthreshold", "Peak test F1\n(diagnostic)"], rows,
               widths=[3.0, 2.1, 2.5, 2.2, 2.4], row_height=0.90, size=17)
    deck.text(slide, "Run 3 default metrics: 1,093 pair TPs / 31.24% F1. Its separate 0.50 sweep records different values.\nThe often-quoted 45.78% is sweep conversion efficiency, not pair recall.", 0.55, 5.15, 12.2, 1.2, 19)


def add_byt5_slides(deck):
    base = "results/stage_byt5_acos/byt5_base_run1"
    dev = read_json(f"{base}/best_dev_metrics.json")
    test = read_json(f"{base}/test_metrics.json")
    settings = read_json(f"{base}/resolved_args.json")
    slide = deck.slide(
        "Phase 6 | generative formulation", "ByT5: One Model for Explicit and Implicit Quads",
        "Byte-level encoder-decoder prediction removes the requirement that every tuple field be an input span.",
        ["src/stage_byt5_acos/model.py", "src/stage_byt5_acos/dataset.py", "src/stage_byt5_acos/linearization.py", github(95)]
    )
    deck.flow(slide, ["Amharic sentence", "UTF-8 bytes", "ByT5 encoder", "Autoregressive decoder"], [
        "Civic / governance text", "256 byte values plus special tokens; no learned subword vocabulary",
        "google/byt5-base; contextual representation", "[ aspect | category | sentiment | opinion ]; can emit NULL"
    ])
    deck.takeaway(slide, "Explicit-only extractors cannot recover NULL quads as implemented. Generation is one solution; the Stage 5 classifier path is another.")

    slide = deck.slide(
        "ByT5 Run 1 | engineering and reproducibility", "The Recorded Training Setup",
        "Resolved run arguments take precedence over current config defaults and earlier slide descriptions.",
        [f"{base}/resolved_args.json", "src/stage_byt5_acos/train.py", github(98, "issues")],
        "The issue reports NaN loss under FP16 and a switch to BF16. Timing and memory reduction claims are issue-level observations, not profiled artifacts reproduced in this audit."
    )
    deck.table(slide, ["Setting", "Recorded value", "Why it matters"], [
        ["Precision / optimizer", f"{settings['precision'].upper()} / {settings['optimizer'].title()}", "Wider exponent range; factorized optimizer state"],
        ["Batch / accumulation", "4 x 4 = 16 effective", "Accumulation controls the effective update size"],
        ["Source / target limit", "768 / 256 byte-token positions", "Finite input/output lengths still truncate"],
        ["Epochs / learning rate", "10 / 1e-4", "Best checkpoint selected on dev"],
        ["Dev selection / batch", "First 1,000 sentences / 16", "Not the full 4,907-sentence dev split"],
    ], widths=[3.5, 3.5, 5.2], row_height=0.59, size=17)
    deck.takeaway(slide, "Reported engineering observations: greedy subset validation was faster than beam search; BF16 avoided the reported FP16 failure.", top=6.05)

    labels = [("Full quadruple", "full_quad"), ("Explicit quadruple", "explicit_quad"), ("Implicit quadruple", "implicit_quad"),
              ("AOPE pair", "aope"), ("ASTE triplet", "aste"), ("ACSE (A,C,S)", "acse"), ("Category + sentiment", "cat_sent")]
    for split, values, subtitle, takeaway in (
        ("Dev", dev, "Best checkpoint (epoch 9 per Issue #98); first 1,000 dev sentences, 1,033 gold quads.",
         "70 implicit true positives are recorded. This is a subset-selection result, not a score on the full dev split."),
        ("Test", test, "Original test split: 4,909 sentences / 5,138 gold quads. Not a cleaned-data evaluation.",
         "811 full-quad true positives, including 274 implicit. Raw counts and split sizes are essential context for these F1 scores."),
    ):
        slide = deck.slide(
            f"ByT5 Run 1 | {split.lower()} results", f"ByT5 {split}: Strict Quadruple and Subtask Metrics", subtitle,
            [f"{base}/{'best_dev_metrics.json' if split == 'Dev' else 'test_metrics.json'}", github(98, "issues")]
        )
        deck.table(slide, ["Task", "Precision", "Recall", "F1", "TP"], [
            [label, percent(values[key]["precision"]), percent(values[key]["recall"]), percent(values[key]["f1"]), f"{values[key]['tp']:,}"]
            for label, key in labels
        ], widths=[4.2, 2.1, 2.1, 2.1, 1.7], row_height=0.47, size=17)
        deck.takeaway(slide, takeaway, top=6.05)

    slide = deck.slide(
        "Implemented after Run 1 | PR #100", "Constrained Decoding and Linguistic Diagnostics",
        "These features are in the codebase; no completed constrained or clitic-normalized ByT5 run is recorded.",
        ["src/stage_byt5_acos/constrained_decoder.py", "src/stage_byt5_acos/evaluate.py", "tests/test_constrained_decoder.py", github(100)],
        "Reproduced with the repository's mock tokenizer: '[ item | PUBLIC_SAFETY#CRIME ' allows only a space, not '|'. The same occurs after completed sentiment plus space. Inline review discussion_r4105131434 identifies this issue."
    )
    deck.columns(slide, "Implemented mechanism", [
        "ByteTrie restricts category and sentiment prefixes during generation.",
        "Evaluator exposes constrained decoding, beam width and generation penalties.",
        "Clitic-normalized matching strips selected affixes from aspect/opinion text; category and sentiment stay exact."
    ], "Validated limitations", [
        "Category/sentiment delimiter transitions can become stuck emitting spaces.",
        "Free-text slots still permit structural bytes; complete grammar validity is not guaranteed.",
        "Repair and test full byte-by-byte paths before claiming zero formatting/category errors."
    ], size=18)

    slide = deck.slide(
        "Implemented after Run 1 | PR #101", "ByT5 Run 2: Implemented, Not Yet Measured",
        "Canonical ordering and regularized token loss are present; curriculum training and scaling remain proposals.",
        ["src/stage_byt5_acos/linearization.py", "src/stage_byt5_acos/train.py", "configs/stage_byt5_acos_run2.yaml", github(101)]
    )
    deck.columns(slide, "Training changes", [
        "Explicit aspects sorted by a_start; implicit aspects with explicit opinions by o_start.",
        "Fully implicit quads placed last, ordered by category and sentiment.",
        "Label-smoothed cross-entropy: epsilon = 0.1. Planned training horizon: 15 epochs."
    ], "Experiment status", [
        "Run 2 standalone evaluation config requests 4 beams and constraints.",
        "Training's internal evaluation does not pass the constrained option; do not assume identical decoding.",
        "Config still points to prepared/, not prepared_clean/. No Run 2 result directory is present."
    ], size=18)


def add_synthesis_slides(deck):
    slide = deck.slide(
        "Cross-architecture synthesis", "Two Evaluation Populations, Not One Leaderboard",
        "Percent F1. The extraction rows use the explicit benchmark; ByT5 uses all original test quadruples.",
        ["results/stage_aope_sdrn/afroxlmr_crf_run1/best_metrics.json", "results/stage_span_aste/", "results/stage_byt5_acos/byt5_base_run1/test_metrics.json"]
    )
    rows = []
    sdrn = read_json("results/stage_aope_sdrn/afroxlmr_crf_run1/best_metrics.json")
    rows.append(["SDRN CRF Run 1", "Explicit pairs", percent(sdrn["f1"]), "Not measured", "Not measured"])
    for label, run in (("Span-ASTE Run 1", "afroxlmr_run1"), ("Span-ASTE Run 2", "afroxlmr_biaffine_run2"), ("Span-ASTE Run 3", "afroxlmr_pooled_focal_run3")):
        result = read_json(f"results/stage_span_aste/{run}/best_metrics.json")
        rows.append([label, "Explicit pairs", percent(result["aope"]["f1"]), percent(result["aste"]["f1"]), "No category head"])
    result = read_json("results/stage_byt5_acos/byt5_base_run1/test_metrics.json")
    rows.append(["ByT5 Run 1", "All quads", percent(result["aope"]["f1"]), percent(result["aste"]["f1"]), percent(result["full_quad"]["f1"])])
    deck.table(slide, ["Model / recorded run", "Population", "AOPE", "ASTE", "Full ACOS"], rows,
               widths=[3.5, 2.3, 1.7, 1.9, 2.8], row_height=0.60, size=16)
    deck.takeaway(slide, "ByT5 records 14.87% implicit-quad F1. A fair ranking requires matching data, tuple fields, implicit coverage and selection protocol.", top=6.05)

    slide = deck.slide(
        "Evaluation protocol", "Exact Match, Subsets and Morphology Must Stay Separate",
        "A credible comparison freezes the data version, checkpoint selection, parser and decoding settings.",
        ["src/pipeline/evaluate.py", "src/stage_byt5_acos/evaluate.py", "src/stage_byt5_acos/linearization.py"]
    )
    deck.columns(slide, "What each score means", [
        "ACOS: (aspect, category, sentiment, opinion). All four must match.",
        "AOPE: (aspect, opinion). ASTE adds sentiment. ACSE omits opinion.",
        "Explicit: neither span is NULL. Implicit: at least one span is NULL.",
        "ByT5 matches text tuples with multiplicity; pipeline metrics use span-index sets."
    ], "What must be controlled", [
        "Do not mix 2,963 unique explicit pairs with 2,983 explicit quads or 5,138 total quads.",
        "Keep strict metrics primary; clitic-normalized scores are a separate diagnostic.",
        "Report malformed-output and truncation rates; parsing can discard invalid outputs.",
        "Use dev for thresholds. Report seeds, uncertainty and held-out sample selection."
    ], size=18)

    slide = deck.slide(
        "Local working-tree artifacts | provisional", "The 500-Sample Files Are Not an LLM Leaderboard",
        "These files were uncommitted at the audit snapshot. Their provenance limits what can be concluded.",
        ["scripts/benchmark_llm_500.py", "scripts/llm_benchmark_full_500.py", "results/llm_test_500_benchmark.json", "results/llm_benchmark_500.json"],
        "No external model API was called in this update. The hybrid benchmark imports 100 fixed predictions and generates the remaining 400 using a keyword lexicon and span heuristics. The other JSON lacks model/version/prompt/request-status metadata."
    )
    provisional = read_json("results/llm_test_500_benchmark.json")
    other = read_json("results/llm_benchmark_500.json")
    deck.table(slide, ["Local artifact", "Samples", "Strict F1", "Interpretation"], [
        ["llm_test_500_benchmark", provisional["samples_evaluated"], percent(provisional["strict_metrics"]["f1"]), "100 fixed predictions + 400 rule outputs"],
        ["llm_benchmark_500", other["samples_evaluated"], percent(other["strict_metrics"]["f1"]), "Only 90 predicted quads; run provenance absent"],
    ], widths=[3.8, 1.3, 1.8, 5.3], row_height=0.90, size=16)
    deck.text(slide, "The hybrid file reports 8.22% clitic-normalized F1 versus 8.02% strict F1.\nThis is not evidence of a ByT5 improvement, independent human agreement, or a controlled external-model comparison.", 0.55, 5.05, 12.2, 1.3, 19)

    slide = deck.slide(
        "Codebase and reproducibility", "What Exists, and What Still Needs Verification",
        "The codebase contains reusable utilities, training/evaluation entry points, configs, metric artifacts and tests.",
        ["src/common/", "src/pipeline/", "tests/", ".github/workflows/ci.yml", "configs/stage_aope_sdrn_clean.yaml"],
        "Static inventory: 12 test files with 88 test functions. CI installs ruff, pytest and pyyaml, not PyTorch/Transformers; optional guards can return without exercising tensor paths. Full ML training and GPU tests were not run for this presentation update."
    )
    deck.columns(slide, "Implemented assets", [
        "common/: span alignment, BIO decoding, label priors and shared PairClassifier.",
        "pipeline/: assembly, inference and explicit/implicit metric views.",
        "SDRN, Span-ASTE and ByT5 each have their own training and evaluation tools.",
        "12 test files / 88 test functions; config and result artifacts document run settings."
    ], "Reproducibility gaps", [
        "Lightweight CI does not exercise all Torch/Transformers paths; test count is not a GPU validation claim.",
        "Some resolved_args files omit the source config path; model checkpoints are not included.",
        "Clean SDRN config reuses the old run output directory: choose a new one before retraining.",
        "No clean-data or ByT5 Run 2 metric artifact is present."
    ], size=17)

    slide = deck.slide(
        "Current roadmap | evidence-ranked", "Data and Evaluation Before Another Scale-Up",
        "Completed implementation is not the same as a completed experiment or a validated research claim.",
        [github(100), github(101), github(103, "issues"), github(104), github(105), "configs/stage_byt5_acos_run2.yaml"]
    )
    deck.table(slide, ["Priority", "Next action", "Evidence of completion needed"], [
        ["1. Data", "Reconcile taxonomy; human-audit a held-out sample", "Adjudicated labels and per-field agreement"],
        ["2. Correctness", "Repair FSM delimiter and structural-byte paths", "Full-sequence generation and parser regression tests"],
        ["3. Experiment", "Run clean-data baselines in fresh output folders", "Same splits, checkpoints and decoding protocol"],
        ["4. Ablation", "Isolate canonical order, smoothing and constraints", "Dev-selected settings; multiple seeds / uncertainty"],
        ["5. Exploration", "Curriculum training or larger ByT5 + LoRA", "Measured gains, resource cost and failure analysis"],
    ], widths=[1.9, 5.3, 5.0], row_height=0.69, size=17)
    deck.text(slide, "Already implemented: data cleaning, word-level SDRN, span pooling, canonical ordering, label smoothing, decoder scaffold and clitic diagnostics.", 0.55, 6.32, 12.2, 0.48, 16, BLUE)

    slide = deck.slide(
        "Coverage and source traceability", "Repository-Wide Review Behind This Update",
        "All 66 PRs and 39 issues were inventoried across open and closed states through #105.",
        ["https://github.com/EthAIResearch/acos-amharic/pull/105", "docs/presentation_v4_audit.md"],
        "Reviewed issue/PR bodies, changed-file inventories, discussion/review records and the working-tree implementation/results. This is a presentation evidence audit, not an exhaustive security or line-by-line correctness audit."
    )
    deck.columns(slide, "Review scope", [
        "65 merged PRs; PR #9 closed unmerged. All 39 issues are closed at the snapshot.",
        "GitHub records include 73 conversation comments, 26 reviews and 43 inline comments.",
        "Source families: common, pipeline, Stages 1-5, SDRN, Span-ASTE and ByT5; data, configs, notebooks and tests.",
        "Original v3.0 is retained unchanged. Source links are included in slide notes."
    ], "Evidence policy", [
        "Prefer actual metric JSON and resolved arguments over narrative summaries.",
        "Identify conflicting artifacts instead of silently combining their numbers.",
        "Do not promote local rule-based outputs to measured LLM performance.",
        "A companion audit maps every issue/PR and lists corrections, additions and verification limits."
    ], size=17)


def github_coverage(number):
    if number in (9, 10, 15, 17, 25, 29, 37, 48, 49, 50, 53, 57, 65, 66, 67, 73, 74, 83, 89, 102, 105):
        return "Supporting configuration, tooling, correction or provenance; reflected in the relevant method/status slides."
    if number <= 10:
        return "Dataset variants and Stage 1 baseline comparison."
    if number <= 12:
        return "Gold-span nearest-distance pairing and error propagation."
    if number <= 22 or number in (35, 36):
        return "Category head, imbalance handling and dual macro-F1 reporting."
    if number in (23, 24, 26, 27, 28, 31, 32, 33):
        return "Sentiment classifier comparison and neutral-class limitations."
    if number <= 45:
        return "Stage 5 implicit detectors and incomplete completion-model coverage."
    if number <= 50:
        return "Pipeline assembly, optional checkpoints and metric views."
    if number <= 80:
        return "SDRN architecture, reweighting/Dice runs, word pooling and seen/unseen diagnostics."
    if number <= 93 or number in (96, 97):
        return "Span-ASTE architecture, corrected metrics and dev-versus-test threshold selection."
    if number <= 99:
        return "ByT5 formulation, recorded training setup and dev/test results."
    if number <= 101:
        return "ByT5 post-Run-1 implementation; unmeasured gains and reproduced FSM limitation."
    return "Clean dataset, taxonomy mismatch and pending clean-data experiments."


def write_audit(entries, output, cache):
    pages = json.loads((cache / "acos-all-issues.json").read_text())
    records = sorted((item for page in pages for item in page), key=lambda item: item["number"])
    pull_pages = json.loads((cache / "acos-all-pulls.json").read_text())
    pulls = {item["number"]: item for page in pull_pages for item in page}
    assert len(records) == 105 and len(pulls) == 66, "GitHub snapshot differs from this audit's scope"
    commit = subprocess.check_output(["git", "rev-parse", "HEAD"], cwd=ROOT, text=True).strip()
    lines = [
        "# Presentation v4.0 Evidence Audit", "", "Audit date: 2026-10-02.", "",
        f"Checkout: `{commit}`. Repository: https://github.com/EthAIResearch/acos-amharic.", "",
        "Original: [Amharic_ACOS_Research_Presentation(v3.0).pptx](Amharic_ACOS_Research_Presentation%28v3.0%29.pptx).",
        f"Updated: [{output.name}]({output.name.replace('(', '%28').replace(')', '%29')}).", "",
        "## Scope", "",
        "All 66 PRs and 39 issues through #105 were reviewed across open and closed states. "
        "65 PRs are merged; #9 was closed unmerged. All 39 issues are closed. "
        "The paginated snapshot includes 73 conversation comments, 26 PR reviews, 43 inline comments "
        "and 351 changed-file entries; no nested review/file connection required another page.", "",
        "The codebase review followed every source family and its configs, tests and result artifacts. "
        "This is a presentation/content audit, not a claim that every line has passed a correctness or security review. "
        "Local uncommitted benchmark files were inspected without modifying or executing their model requests.", "",
        "## Important Corrections", "",
        "- Domain is civic/governance commentary, not e-commerce/hospitality.",
        "- 22 declared categories, 23 observed across splits; one test-only ECONOMY#UTILITIES annotation is outside the allowed list. Neither 13 nor 29 is the current declared count.",
        "- Original train/dev/test sizes are 39,261/4,907/4,909 sentences and 41,224/5,129/5,138 quads. Implicitness is 42.71% in train but 41.94% in original test.",
        "- Clean train/dev/test sizes are 38,385/4,892/4,908 sentences and 40,426/5,110/5,133 quads. No existing model score is relabeled as a clean-data result.",
        "- Stage 5 contains learned implicit detectors and optional category/sentiment paths. Generation is not the only possible NULL-handling method.",
        "- Do not multiply component F1 scores to derive end-to-end recall. Issue #51 reports approximately 14% pipeline F1; no standalone pipeline metric JSON was found.",
        "- Span-ASTE default result files give Run 1/2/3 AOPE F1 of 30.80%/32.82%/31.24%, OTE F1 of 39.48%/42.87%/42.89%, and OTE recall of 30.52%/35.57%/46.73%.",
        "- Run 2's 34.02% and Run 3 sweep's 45.78% are candidate conversion rates, not AOPE F1 or pair recall. Run 3 best_metrics conversion is 45.04%; separate sweep/default artifacts disagree and remain labeled separately.",
        "- Span-ASTE Run 2 dev-best AOPE threshold is 0.50, with 32.90% test sweep F1. Its 33.31% at 0.35 is a test-selected diagnostic peak. Run 3 dev-best threshold is 0.65, with 31.97% test sweep F1; 32.04% at 0.60 is test-selected.",
        "- Run 3 enabled mean pooling and weighted cross-entropy, with weights [0.1, 1.0, 1.0, 2.0]. use_focal_loss is false. These changes were not isolated ablations.",
        "- Prefix pooling uses half-open [s,e) spans: (P[e]-P[s])/(e-s), after an O(n) prefix pass. Endpoints are contextual representations, so saying they completely ignore internal words is too strong.",
        "- 55.20% aspect F1 is best among the three Span-ASTE runs, not repository-wide: SDRN CRF Run 1 records 60.41%.",
        "- ByT5 Run 1 effective batch is 4 x 4 = 16, source/target limits are 768/256, and best-dev metrics cover the first 1,000 sentences (1,033 gold quads), not full dev.",
        "- ACSE F1 is not ATE F1. Explicit-only AOPE/ASTE and all-quad ByT5 scores are not a directly comparable leaderboard.",
        "- Constrained decoding and clitic metrics are implemented in PR #100. Canonical ordering and smoothing are implemented in #101. No corresponding completed ByT5 result is recorded.",
        "- Reproduced FSM delimiter stall after a complete category/sentiment followed by space. The required pipe is disallowed. Free-text slots also allow structural bytes. Do not claim guaranteed valid outputs.",
        "- ByT5 Run 2 still uses prepared/. Training's internal evaluator does not enable the configured standalone FSM. Clean SDRN config uses clean paths but reuses an existing output directory.",
        "- Local 500-sample hybrid result is 8.02% strict / 8.22% normalized F1, from 100 fixed predictions plus 400 keyword-rule outputs. The other local file records 1.96% strict F1 with 90 predictions over 500 sentences and lacks model/run provenance. Neither is a controlled LLM leaderboard entry.",
        "- Unseen phrases are not tokenizer unknown tokens. Byte representation coverage does not prove linguistic generalization. Claims of inherent BIO incompatibility and universally necessary generation were removed.", "",
        "## Newly Added Evidence", "",
        "Dataset provenance, split sizes and authentic Amharic examples; duplicate/leakage cleanup; schema mismatch; Stage 1/3/4 backbone comparisons; "
        "Stage 5 detector results; SDRN reweighting, Dice and word-level history; seen/unseen opinion diagnostics; "
        "Run 2 status; metric definitions; local benchmark caveats; code/CI limitations; prioritized validation roadmap.", "",
        "## Verification and Limits", "",
        "- Original PowerPoint hash checked before and after each build; original PDF also left unchanged.",
        "- Generated deck reopened; slide count and unique ZIP package parts checked.",
        "- Dataset counts, declared/observed labels, normalized-text duplicates and pairwise split overlap recomputed from JSONL.",
        "- Clean data independently checked for zero normalized-text duplicates, overlap, duplicate quads and same-span sentiment conflicts. This does not establish human annotation correctness or absence of semantic/near-duplicate leakage.",
        "- Decoder limitation reproduced with the repository's MockByT5Tokenizer. No ML source files were changed to repair it.",
        "- Static test inventory: 12 files, 88 test functions. Full ML tests, training and GPU evaluation were not run. CI's minimal dependencies and optional import guards do not establish tensor-path coverage.",
        "- No external LLM requests, training jobs, GitHub edits, commits or branch changes were performed.",
        "- Final visual and PDF checks are performed after generation; see the delivered PDF for the rendered artifact.", "",
        "## Slide-Level Sources", "",
    ]
    for number, entry in enumerate(entries, 1):
        lines.extend([f"### {number}. {entry['title']}", ""])
        if entry["note"]:
            lines.extend([entry["note"], ""])
        for source in entry["sources"]:
            target = source if source.startswith("https://") else "../" + source
            lines.append(f"- [{source}]({target})")
        lines.append("")
    lines.extend(["## Every Issue and PR", "", "Low-level fixes and tooling changes are grouped into the relevant method/provenance section rather than given one slide each.", "",
                  "| Record | State | Title | Presentation disposition |", "| --- | --- | --- | --- |"])
    for record in records:
        number = record["number"]
        state = "merged" if number in pulls and pulls[number]["merged_at"] else record["state"]
        kind = "PR" if number in pulls else "Issue"
        title = record["title"].replace("|", "\\|").replace("\n", " ")
        lines.append(f"| [{kind} #{number}]({record['html_url']}) | {state} | {title} | {github_coverage(number)} |")
    lines.extend(["", "## Regeneration", "", "The PowerPoint builder requires python-pptx; PDF rendering requires LibreOffice. Presentation tools are kept separate from ML requirements.", "",
                  "```bash", "python scripts/update_research_presentation.py --build", "libreoffice --headless --convert-to pdf --outdir docs 'docs/Amharic_ACOS_Research_Presentation(v4.0).pptx'", "```", "",
                  "To regenerate this GitHub ledger, supply `--github-cache DIR` containing the paginated `gh api --paginate --slurp` issue and pull-list JSON snapshots named `acos-all-issues.json` and `acos-all-pulls.json`. The script validates the audited 105-record scope.", ""])
    audit = ROOT / "docs/presentation_v4_audit.md"
    audit.write_text("\n".join(lines), encoding="utf-8")
    print(f"Wrote evidence ledger: {audit}")


def build_presentation(source, output):
    original_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    deck = Deck(source)
    add_problem_slide(deck)
    add_dataset_slides(deck)
    add_example_slide(deck)
    add_pipeline_slides(deck)
    add_sdrn_slides(deck)
    add_span_slides(deck)
    add_byt5_slides(deck)
    add_synthesis_slides(deck)
    deck.save(output)
    assert hashlib.sha256(source.read_bytes()).hexdigest() == original_hash
    expected = len(deck.entries)
    assert len(slide_texts(output)) == expected
    print(f"Wrote {expected} slides: {output}")
    return deck.entries


def slide_texts(path):
    with zipfile.ZipFile(path) as archive:
        slide_names = sorted(
            (
                name for name in archive.namelist()
                if name.startswith("ppt/slides/slide") and name.endswith(".xml")
            ),
            key=lambda name: int(name.rsplit("slide", 1)[1][:-4]),
        )
        return [
            "\n".join(
                node.text or ""
                for node in ET.fromstring(archive.read(name)).findall(".//a:t", DRAWING_NS)
            )
            for name in slide_names
        ]


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=SOURCE)
    parser.add_argument("--output", type=Path, default=OUTPUT)
    parser.add_argument("--build", action="store_true")
    parser.add_argument("--github-cache", type=Path, help="Directory containing paginated GitHub JSON snapshots for the audit ledger")
    args = parser.parse_args()
    if args.build:
        if args.source.resolve() == args.output.resolve():
            parser.error("The output must not overwrite the original presentation.")
        entries = build_presentation(args.source, args.output)
        if args.github_cache:
            write_audit(entries, args.output, args.github_cache)
        return
    texts = slide_texts(args.source)
    print(f"{len(texts)} slides in {args.source.name}")
    for number, text in enumerate(texts, 1):
        print(f"{number:02}: {' | '.join(text.splitlines()[:2])}")


if __name__ == "__main__":
    main()